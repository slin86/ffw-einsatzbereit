import csv
import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from einsatzbereit.services.status import Cell, CellStatus, Matrix

STATUS_LABEL: dict[CellStatus, str] = {
    CellStatus.MISSING: "fehlt",
    CellStatus.EXPIRED: "abgelaufen",
    CellStatus.EXPIRING: "läuft ab",
    CellStatus.VALID: "gültig",
    CellStatus.NOT_REQUIRED: "",
}

STATUS_COLOR: dict[CellStatus, str] = {
    CellStatus.MISSING: "E8B4AE",
    CellStatus.EXPIRED: "E8B4AE",
    CellStatus.EXPIRING: "F3D98B",
    CellStatus.VALID: "B9DCC4",
    CellStatus.NOT_REQUIRED: "FFFFFF",
}


def _fmt(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def cell_text(cell: Cell) -> str:
    """
    Returns the text of one matrix cell. Missing entries show fehlt, unlimited ones show
    unbefristet, and certifications that are not required show their expiry date in brackets.
    """
    if cell.status == CellStatus.MISSING:
        return "fehlt"
    if cell.status == CellStatus.NOT_REQUIRED:
        return f"({_fmt(cell.expires_on)})" if cell.expires_on else ""
    if cell.expires_on is None:
        return "unbefristet"
    return _fmt(cell.expires_on)


def _header(matrix: Matrix) -> list[str]:
    return ["Nr", "Name", "Vorname", *[c.name for c in matrix.certifications]]


def to_csv(matrix: Matrix) -> bytes:
    """
    Writes the matrix as semicolon separated CSV with a status column and an expiry column per
    certification. A byte order mark lets Excel detect UTF8.
    """
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    header = ["Nr", "Name", "Vorname"]
    for c in matrix.certifications:
        header += [f"{c.name} Status", f"{c.name} gültig bis"]
    writer.writerow(header)
    for row in matrix.rows:
        line = [row.member.number, row.member.last_name, row.member.first_name]
        for cell in row.cells:
            line += [STATUS_LABEL[cell.status], _fmt(cell.expires_on)]
        writer.writerow(line)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx(matrix: Matrix, filter_description: str) -> bytes:
    """
    Builds an Excel sheet with one row per member, one colored cell per certification, a frozen
    header and a color legend.
    """
    wb = Workbook()
    wb.remove(wb["Sheet"])
    ws = wb.create_sheet("Nachweise")
    ws.append([f"Einsatzbereit – Stand {_fmt(matrix.today)}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([filter_description])
    ws.append([])
    ws.append(_header(matrix))
    header_row = ws.max_row
    for head_cell in ws[header_row]:
        head_cell.font = Font(bold=True)
        head_cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    for row in matrix.rows:
        ws.append(
            [row.member.number, row.member.last_name, row.member.first_name]
            + [cell_text(c) for c in row.cells]
        )
        for idx, status_cell in enumerate(row.cells, start=4):
            xl = ws.cell(row=ws.max_row, column=idx)
            xl.fill = PatternFill("solid", fgColor=STATUS_COLOR[status_cell.status])
            xl.alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"D{header_row + 1}"
    for col in range(1, len(matrix.certifications) + 4):
        ws.column_dimensions[get_column_letter(col)].width = 8 if col == 1 else 16
    ws.append([])
    ws.append(["Legende:", "gültig", "läuft ab", "fehlt/abgelaufen"])
    for col, status in ((2, CellStatus.VALID), (3, CellStatus.EXPIRING), (4, CellStatus.EXPIRED)):
        ws.cell(row=ws.max_row, column=col).fill = PatternFill(
            "solid", fgColor=STATUS_COLOR[status]
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(matrix: Matrix, filter_description: str) -> bytes:
    """
    Builds a landscape PDF of the matrix. Columns use the certification abbreviations to fit the
    page, and a legend below the table explains colors and abbreviations.
    """
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Einsatzbereit – Nachweise",
    )
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7, leading=8)
    head = ParagraphStyle("head", parent=small, fontName="Helvetica-Bold")

    header = ["Nr", "Name", "Vorname", *[c.short_name or c.name for c in matrix.certifications]]
    data: list[list[Paragraph]] = [[Paragraph(h, head) for h in header]]
    style: list[Any] = [
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9AA5B1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE3EA")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("REPEATROWS", (0, 0), (-1, 0)),
    ]
    for r, row in enumerate(matrix.rows, start=1):
        data.append(
            [
                Paragraph(row.member.number, small),
                Paragraph(row.member.last_name, small),
                Paragraph(row.member.first_name, small),
            ]
            + [Paragraph(cell_text(c), small) for c in row.cells]
        )
        for c_idx, cell in enumerate(row.cells, start=3):
            style.append(
                (
                    "BACKGROUND",
                    (c_idx, r),
                    (c_idx, r),
                    colors.HexColor("#" + STATUS_COLOR[cell.status]),
                )
            )

    n_certs = max(len(matrix.certifications), 1)
    available = landscape(A4)[0] - 20 * mm
    fixed = [14 * mm, 32 * mm, 28 * mm]
    cert_width = (available - sum(fixed)) / n_certs
    table = Table(data, colWidths=fixed + [cert_width] * len(matrix.certifications), repeatRows=1)
    table.setStyle(TableStyle(style))

    story = [
        Paragraph(f"Einsatzbereit – Stand {_fmt(matrix.today)}", styles["Title"]),
        Paragraph(filter_description, styles["Normal"]),
        Spacer(1, 4 * mm),
        table,
        Spacer(1, 4 * mm),
        Paragraph(
            "Grün: gültig · Gelb: läuft bald ab · Rot: fehlt oder abgelaufen · "
            "Werte in Klammern: nicht gefordert",
            small,
        ),
        Paragraph(
            " · ".join(f"{c.short_name} = {c.name}" for c in matrix.certifications if c.short_name),
            small,
        ),
    ]
    doc.build(story)
    return out.getvalue()
